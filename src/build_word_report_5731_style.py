from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parents[1]
OUT_DIR = BASE_DIR / "outputs" / "05_laporan_word"
OUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT_DIR / "Laporan_UTS_Battery_Health_Decision_Tree_Format_5731.docx"
FLOW_PATH = OUT_DIR / "alur_penelitian_format_5731.png"

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def set_run(run, size=12, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = BLACK


def add_paragraph(doc, text="", size=12, bold=False, italic=False, align=None, after=6, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.38)
    if align is not None:
        p.alignment = align
        p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run(run, size=12, bold=True)
    return p


def add_code_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    set_run(run, size=12, bold=True)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Inches(0)
    for line in code.strip("\n").splitlines():
        run = p.add_run(line.rstrip() + "\n")
        set_run(run, size=9.5, font="Consolas")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run(run, size=11, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run(run, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_picture(doc, path, width, caption):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_after = Pt(10)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run(r, size=11, bold=False)


def make_flowchart():
    img = Image.new("RGB", (1400, 780), "white")
    draw = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 36)
        box_font = ImageFont.truetype("arialbd.ttf", 22)
        small_font = ImageFont.truetype("arial.ttf", 18)
    except OSError:
        title_font = box_font = small_font = ImageFont.load_default()

    draw.text((70, 35), "Diagram Alur Penelitian", fill="#000000", font=title_font)
    steps = [
        ("Pengumpulan NASA Battery Dataset", "Dataset ZIP dan file MAT"),
        ("Ekstraksi Data Discharge", "Pembentukan battery_features.csv"),
        ("EDA", "Distribusi kelas, SOH, korelasi"),
        ("Preprocessing", "Missing value, outlier, seleksi fitur"),
        ("Pemodelan Decision Tree", "GridSearchCV dan validasi grup"),
        ("Evaluasi Model", "Akurasi, precision, recall, F1"),
        ("Simpan Model", "decision_tree_battery.pkl"),
        ("Deployment Web", "Dashboard Streamlit"),
    ]
    xy = [(70, 130), (530, 130), (990, 130), (70, 330), (530, 330), (990, 330), (300, 560), (760, 560)]
    w, h = 330, 105
    for (title, note), (x, y) in zip(steps, xy):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=12, fill="#f7f7f7", outline="#111111", width=2)
        draw.text((x + 18, y + 22), title, fill="#000000", font=box_font)
        draw.text((x + 18, y + 58), note, fill="#333333", font=small_font)
    for a, b in [
        ((400, 182), (520, 182)), ((860, 182), (980, 182)), ((1155, 235), (1155, 320)),
        ((400, 382), (520, 382)), ((860, 382), (980, 382)), ((1120, 435), (930, 555)),
        ((630, 612), (750, 612)),
    ]:
        draw.line((a, b), fill="#000000", width=3)
        ex, ey = b
        draw.polygon([(ex, ey), (ex - 14, ey - 8), (ex - 14, ey + 8)], fill="#000000")
    img.save(FLOW_PATH)


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.1)
    sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)


def cover(doc):
    add_paragraph(doc, "Prediksi Kondisi Kesehatan Baterai Lithium-Ion", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first_line=False)
    add_paragraph(doc, "Menggunakan Decision Tree pada NASA Battery Dataset", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=80, first_line=False)
    add_paragraph(doc, "disusun oleh", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, first_line=False)
    add_paragraph(doc, "[Nama Mahasiswa]", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first_line=False)
    add_paragraph(doc, "[NIM]", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=150, first_line=False)
    add_paragraph(doc, "FAKULTAS ILMU KOMPUTER", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first_line=False)
    add_paragraph(doc, "UNIVERSITAS AMIKOM YOGYAKARTA", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first_line=False)
    add_paragraph(doc, "YOGYAKARTA", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first_line=False)
    add_paragraph(doc, "2026", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, first_line=False)
    doc.add_page_break()


def build():
    make_flowchart()
    doc = Document()
    configure(doc)
    cover(doc)

    add_heading(doc, "1. LATAR BELAKANG PENELITIAN")
    for text in [
        "Baterai lithium-ion merupakan salah satu komponen energi yang banyak digunakan pada perangkat elektronik, kendaraan listrik, sistem industri, dan penyimpanan energi. Keunggulan baterai ini terletak pada densitas energi yang tinggi dan umur pakai yang relatif panjang. Meskipun demikian, penggunaan berulang akan menyebabkan kapasitas baterai menurun secara bertahap. Penurunan kapasitas tersebut dapat memengaruhi performa perangkat dan meningkatkan risiko gangguan operasional.",
        "Kondisi kesehatan baterai umumnya dinilai melalui State of Health (SOH), yaitu perbandingan antara kapasitas aktual dengan kapasitas awal atau nominal. Apabila SOH terus menurun, baterai dapat masuk ke kondisi Warning atau End of Life. Oleh karena itu, diperlukan metode yang mampu membantu mengklasifikasikan kondisi baterai berdasarkan data pengukuran sensor.",
        "Dataset yang digunakan dalam penelitian ini adalah NASA Battery Dataset. Dataset tersebut berisi data eksperimen baterai lithium-ion dalam proses charge, discharge, dan impedance. Pada penelitian ini, fokus diarahkan pada data discharge karena data tersebut memuat informasi kapasitas dan pola sensor yang relevan untuk melihat degradasi baterai.",
        "Model yang digunakan adalah Decision Tree. Model ini dipilih karena hasilnya lebih mudah dijelaskan dibandingkan beberapa model machine learning lain. Struktur pohon keputusan dapat dianalisis melalui aturan keputusan sehingga sesuai untuk laporan akademik. Namun, Decision Tree memiliki risiko overfitting apabila kedalaman pohon tidak dikontrol. Oleh karena itu, penelitian ini menggunakan pembatasan hyperparameter, GridSearchCV, dan validasi berbasis battery_id.",
        "Klaim hasil penelitian ini dibatasi pada data hasil ekstraksi NASA Battery Dataset yang digunakan dalam proyek. Model tidak dimaksudkan untuk langsung digeneralisasi ke semua jenis baterai lithium-ion di dunia nyata tanpa pengujian tambahan, karena karakteristik degradasi baterai sangat dipengaruhi oleh kondisi operasional, desain sel, suhu, dan pola penggunaan.",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "2. DIAGRAM ALUR PENELITIAN")
    add_paragraph(doc, "Diagram alur penelitian ditampilkan pada Gambar 1 berikut ini. Diagram tersebut menunjukkan tahapan penelitian mulai dari pengumpulan dataset hingga deployment aplikasi web.")
    add_picture(doc, FLOW_PATH, 6.2, "Gambar 1. Diagram Alur Penelitian")
    add_paragraph(doc, "Alur penelitian dimulai dari pengumpulan NASA Battery Dataset dalam format ZIP dan MAT. Data kemudian diekstraksi untuk mengambil siklus discharge dan membentuk dataset fitur battery_features.csv. Setelah dataset terbentuk, dilakukan EDA untuk memahami distribusi kelas, pola SOH, dan hubungan antar fitur.")
    add_paragraph(doc, "Tahap berikutnya adalah preprocessing dan feature engineering. Pada tahap ini, fitur yang terlalu dekat dengan target seperti capacity_ahr dan soh tidak digunakan sebagai input model. Model Decision Tree kemudian dilatih dan dievaluasi menggunakan nested StratifiedGroupKFold berdasarkan battery_id. Setelah model akhir diperoleh, model disimpan dalam format pkl dan digunakan pada dashboard web berbasis Streamlit.")

    add_heading(doc, "3. LAMPIRAN DAN ANALISIS KODE PROGRAM")
    add_heading(doc, "3.1 Pengambilan Data")
    add_paragraph(doc, "Tahap pengambilan data dilakukan melalui pembentukan dataset fitur dari data NASA Battery Dataset. File utama untuk membangun dataset adalah src/build_dataset.py, sedangkan dataset hasil olahan disimpan pada data/battery_features.csv. Dataset ini berisi 2.794 baris discharge cycle dari 34 baterai unik.")
    add_code_title(doc, "Inventarisasi Dataset dan Pembacaan Data")
    add_code(doc, """
df = pd.read_csv(DATA_PATH)
df = df.drop_duplicates()
df = df.dropna(subset=[TARGET, GROUP_COLUMN])

drop_columns = [
    GROUP_COLUMN,
    TARGET,
    "capacity_ahr",
    "soh",
    "elapsed_time_sec",
    "cycle_index",
    "discharge_index",
]
feature_columns = [col for col in df.columns if col not in drop_columns]
X = df[feature_columns]
y = df[TARGET]
groups = df[GROUP_COLUMN]
""")
    add_paragraph(doc, "Kode tersebut membaca dataset hasil ekstraksi, menghapus data duplikat, memastikan target dan group tersedia, serta memisahkan kolom fitur dan target. Kolom capacity_ahr dan soh tidak digunakan sebagai fitur karena keduanya berkaitan langsung dengan pembentukan label health_status.")

    add_heading(doc, "3.2 Exploratory Data Analysis (EDA)")
    add_paragraph(doc, "EDA dilakukan untuk memahami karakteristik dataset sebelum proses pemodelan. Analisis ini mencakup distribusi kelas health_status, tren SOH, dan korelasi antar fitur numerik. Hasil EDA disimpan pada folder outputs/01_eda.")
    add_picture(doc, BASE_DIR / "outputs" / "01_eda" / "eda_class_distribution.png", 4.9, "Gambar 2. Distribusi Kelas Health Status")
    add_picture(doc, BASE_DIR / "outputs" / "01_eda" / "eda_soh_trend.png", 5.8, "Gambar 3. Tren SOH Baterai")
    add_picture(doc, BASE_DIR / "outputs" / "01_eda" / "eda_correlation_heatmap.png", 5.8, "Gambar 4. Heatmap Korelasi Fitur")
    add_paragraph(doc, "Visualisasi distribusi kelas digunakan untuk melihat keseimbangan data antar kelas. Tren SOH menunjukkan proses penurunan kesehatan baterai secara bertahap. Heatmap korelasi digunakan untuk melihat hubungan antar fitur sensor sehingga dapat membantu memahami pola data sebelum model dilatih.")

    add_heading(doc, "3.3 Preprocessing")
    add_paragraph(doc, "Preprocessing bertujuan membuat data lebih siap digunakan oleh model. Pada penelitian ini, missing value ditangani menggunakan median imputer, sedangkan outlier ditangani menggunakan IQR capping. Scaling tidak menjadi tahap utama karena Decision Tree tidak sensitif terhadap perbedaan skala fitur.")
    add_code_title(doc, "Pipeline Preprocessing dan Model")
    add_code(doc, """
pipeline = Pipeline(
    steps=[
        ("feature_engineering", BatteryFeatureEngineer()),
        ("outlier_capping", IQRClipper(factor=1.5)),
        ("imputer", SimpleImputer(strategy="median")),
        ("model", DecisionTreeClassifier(random_state=42, class_weight="balanced")),
    ]
)
""")
    add_paragraph(doc, "Pipeline tersebut memastikan seluruh proses preprocessing dilakukan secara konsisten di setiap fold validasi. Dengan cara ini, proses imputasi dan outlier capping hanya dipelajari dari data latih pada masing-masing fold, sehingga risiko kebocoran data dapat dikurangi.")

    add_heading(doc, "3.4 Feature Engineering")
    add_paragraph(doc, "Feature engineering dilakukan secara konservatif. Fitur turunan yang digunakan berasal dari pola tegangan, arus, dan suhu, seperti voltage_range, current_range, temperature_range, voltage_stability_ratio, dan temperature_stability_ratio. Fitur yang terlalu dekat dengan target tidak digunakan agar model tidak terlihat bagus hanya karena shortcut prediksi.")
    add_table(doc, ["Kelompok Fitur", "Contoh Fitur", "Fungsi"], [
        ("Tegangan", "voltage_mean, voltage_min, voltage_std", "Menggambarkan karakteristik tegangan saat discharge"),
        ("Arus", "current_mean, current_min, current_std", "Menggambarkan pola beban dan kestabilan arus"),
        ("Suhu", "temperature_mean, temperature_min, temperature_std", "Menggambarkan respons termal baterai"),
        ("Lingkungan", "ambient_temperature", "Menunjukkan kondisi suhu eksperimen"),
    ], [1.5, 2.5, 2.5])

    add_heading(doc, "3.5 Pemodelan Decision Tree")
    add_paragraph(doc, "Model Decision Tree dilatih dengan hyperparameter tuning menggunakan GridSearchCV. Parameter yang diuji meliputi criterion, max_depth, min_samples_split, dan min_samples_leaf. Metrik yang digunakan untuk tuning adalah f1_weighted karena target terdiri dari tiga kelas dan distribusinya tidak sepenuhnya seimbang.")
    add_code_title(doc, "Hyperparameter Tuning Decision Tree")
    add_code(doc, """
param_grid = {
    "model__criterion": ["gini", "entropy"],
    "model__max_depth": [3, 4, 5, 6, 7, 8],
    "model__min_samples_split": [2, 5, 10],
    "model__min_samples_leaf": [1, 3, 5, 10],
}

grid = GridSearchCV(
    estimator=pipeline,
    param_grid=param_grid,
    scoring="f1_weighted",
    cv=inner_cv,
)
""")
    add_paragraph(doc, "Parameter terbaik model akhir adalah criterion gini, max_depth 6, min_samples_leaf 1, dan min_samples_split 5. Pembatasan max_depth digunakan untuk membantu mengurangi overfitting pada model Decision Tree.")
    add_picture(doc, BASE_DIR / "outputs" / "03_modeling" / "decision_tree.png", 6.1, "Gambar 5. Struktur Decision Tree Model Akhir")

    add_heading(doc, "3.6 Evaluasi Model")
    add_paragraph(doc, "Evaluasi model dilakukan menggunakan nested StratifiedGroupKFold berdasarkan battery_id. Strategi ini lebih aman dibanding split acak biasa karena data dari baterai yang sama tidak dicampur secara tidak tepat antara data latih dan data uji. Dengan demikian, model diuji pada baterai yang berbeda dari data latih.")
    add_table(doc, ["Metrik", "Nilai"], [
        ("Akurasi", "0,8446 +/- 0,0408"),
        ("F1-score berbobot", "0,8491 +/- 0,0445"),
        ("Precision berbobot", "0,8712 +/- 0,0558"),
        ("Recall berbobot", "0,8446 +/- 0,0408"),
        ("Selisih F1 train-test", "0,0976 +/- 0,0539"),
    ], [2.4, 4.1])
    add_picture(doc, BASE_DIR / "outputs" / "04_evaluation" / "confusion_matrix.png", 4.8, "Gambar 6. Confusion Matrix")
    add_picture(doc, BASE_DIR / "outputs" / "04_evaluation" / "feature_importance.png", 5.7, "Gambar 7. Feature Importance")
    add_paragraph(doc, "Berdasarkan classification report out-of-fold, kelas Healthy dan End_of_Life memiliki F1-score yang lebih tinggi dibanding kelas Warning. Kelas Warning lebih sulit diprediksi karena berada pada area transisi antara kondisi baterai sehat dan kondisi baterai yang mendekati akhir masa pakai.")

    add_heading(doc, "3.7 Ringkasan Hasil Evaluasi")
    add_table(doc, ["Kelas", "Precision", "Recall", "F1-score", "Support"], [
        ("End_of_Life", "0,90", "0,87", "0,88", "1318"),
        ("Healthy", "0,92", "0,84", "0,88", "938"),
        ("Warning", "0,60", "0,74", "0,66", "538"),
        ("Weighted Avg", "0,85", "0,83", "0,84", "2794"),
    ], [1.7, 1.2, 1.2, 1.2, 1.2])
    add_paragraph(doc, "Hasil tersebut menunjukkan bahwa model cukup baik untuk membedakan kondisi baterai yang sudah jelas, tetapi masih memiliki tantangan pada kelas Warning. Hal ini wajar karena kelas Warning merupakan kondisi peralihan yang secara sensor dapat memiliki kemiripan dengan Healthy maupun End_of_Life.")

    add_heading(doc, "4. LAMPIRAN DAN ANALISIS DEPLOYMENT WEB")
    add_heading(doc, "4.1 Deskripsi Singkat Deployment")
    add_paragraph(doc, "Model Decision Tree yang telah dilatih dan dievaluasi di-deploy ke dalam aplikasi web berbasis Streamlit. Aplikasi ini dibuat pada file app.py dan memuat model dari models/decision_tree_battery.pkl. Pengguna dapat memasukkan nilai sensor discharge melalui sidebar, kemudian sistem menampilkan hasil prediksi kondisi baterai.")
    add_code_title(doc, "Menjalankan Aplikasi Web")
    add_code(doc, "streamlit run app.py")

    add_heading(doc, "4.2 Isi Dashboard Web")
    for item in [
        "Dashboard menampilkan ringkasan model, akurasi, F1-score berbobot, dan jumlah fitur.",
        "Sidebar digunakan sebagai panel input nilai sensor baterai.",
        "Bagian hasil prediksi menampilkan kelas Healthy, Warning, atau End_of_Life.",
        "Grafik probabilitas kelas menampilkan tingkat keyakinan model terhadap setiap kelas.",
        "Grafik feature importance dan performa per fold membantu pengguna memahami hasil model.",
    ]:
        add_paragraph(doc, item, first_line=False)

    add_heading(doc, "4.3 Screenshot Dashboard")
    add_paragraph(doc, "Pada bagian ini, lampirkan screenshot aplikasi Streamlit yang telah dijalankan melalui browser. Screenshot minimal yang disarankan adalah tampilan dashboard utama, input prediksi pada sidebar, hasil prediksi, dan grafik evaluasi model.")
    add_paragraph(doc, "[Tempel Screenshot Deployment 1: Tampilan dashboard utama]", italic=True, first_line=False)
    add_paragraph(doc, "[Tempel Screenshot Deployment 2: Sidebar input prediksi]", italic=True, first_line=False)
    add_paragraph(doc, "[Tempel Screenshot Deployment 3: Hasil prediksi dan probabilitas kelas]", italic=True, first_line=False)

    add_heading(doc, "4.4 Analisis Singkat Deployment")
    add_paragraph(doc, "Deployment menggunakan Streamlit memudahkan pengguna untuk mencoba model tanpa harus menjalankan kode training. Aplikasi web ini cocok sebagai prototype akademik karena dapat menampilkan input, output prediksi, dan visualisasi model dalam satu halaman. Meskipun demikian, aplikasi masih berjalan secara lokal sehingga untuk penggunaan luas perlu dilakukan deployment ke layanan cloud.")

    add_heading(doc, "5. LAMPIRAN")
    for item in [
        "Dataset hasil olahan: data/battery_features.csv",
        "Kode training dan evaluasi: train_decision_tree.py",
        "Kode komponen pipeline: model_components.py",
        "Kode aplikasi web: app.py",
        "Model akhir: models/decision_tree_battery.pkl",
        "Output EDA: outputs/01_eda",
        "Output modelling: outputs/03_modeling",
        "Output evaluasi: outputs/04_evaluation",
        "Repository GitHub: https://github.com/FikihRizaldi/Battery-Health-Decision-Tree.git",
    ]:
        add_paragraph(doc, item, first_line=False)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
