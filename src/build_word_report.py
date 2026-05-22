from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "outputs" / "05_laporan_word"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUTPUT_DIR / "Laporan_UTS_Battery_Health_Decision_Tree.docx"
FLOWCHART_PATH = OUTPUT_DIR / "alur_penelitian.png"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 31, 46)
MUTED = RGBColor(90, 103, 120)
LIGHT_FILL = "F2F4F7"
MID_FILL = "E8EEF5"
WHITE = RGBColor(255, 255, 255)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=11, bold=False, italic=False, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_para(doc, text="", style=None, size=11, bold=False, italic=False, color=None, align=None, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color or INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level <= 2 else DARK_BLUE)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11, color=INK)
    return p


def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    for line in code.strip().splitlines():
        run = p.add_run(line.rstrip() + "\n")
        set_run(run, size=9.2, color=RGBColor(35, 45, 60), font="Consolas")
    p.paragraph_format.left_indent = Inches(0.18)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run(run, size=9.5, italic=True, color=MUTED)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run(run, size=10.5, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx != len(row) - 1 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run(run, size=10.2, color=INK)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_fill(cell, MID_FILL)
    p = cell.paragraphs[0]
    run = p.add_run(title + "\n")
    set_run(run, size=10.8, bold=True, color=DARK_BLUE)
    run = p.add_run(text)
    set_run(run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_picture_if_exists(doc, path, width, caption):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)


def create_flowchart():
    width, height = 1600, 900
    img = Image.new("RGB", (width, height), "#f6f8fb")
    draw = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 42)
        node_font = ImageFont.truetype("arialbd.ttf", 24)
        small_font = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        title_font = node_font = small_font = ImageFont.load_default()

    draw.text((70, 48), "Alur Penelitian Proyek Data Mining", fill="#1f4d78", font=title_font)
    steps = [
        ("01", "Dataset NASA Battery", "Data mentah baterai lithium-ion"),
        ("02", "Pembentukan Dataset", "Ekstraksi fitur discharge"),
        ("03", "EDA", "Distribusi kelas dan pola sensor"),
        ("04", "Preprocessing", "Seleksi fitur, imputasi, outlier capping"),
        ("05", "Training Decision Tree", "Tuning parameter model"),
        ("06", "Validasi Model", "Nested StratifiedGroupKFold"),
        ("07", "Evaluasi", "Akurasi, precision, recall, F1-score"),
        ("08", "Deployment", "Dashboard prediksi Streamlit"),
    ]
    positions = [
        (70, 150), (470, 150), (870, 150), (70, 350),
        (470, 350), (870, 350), (270, 580), (700, 580),
    ]
    box_w, box_h = 320, 120
    for (num, title, note), (x, y) in zip(steps, positions):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=18, fill="#ffffff", outline="#d7dfe8", width=3)
        draw.text((x + 24, y + 18), f"TAHAP {num}", fill="#2e74b5", font=small_font)
        draw.text((x + 24, y + 50), title, fill="#141f2e", font=node_font)
        draw.text((x + 24, y + 84), note, fill="#5a6778", font=small_font)
    arrows = [
        ((390, 210), (460, 210)), ((790, 210), (860, 210)),
        ((1030, 270), (1030, 340)), ((390, 410), (460, 410)),
        ((790, 410), (860, 410)), ((630, 470), (520, 570)),
        ((590, 640), (690, 640)),
    ]
    for start, end in arrows:
        draw.line((start, end), fill="#2e74b5", width=5)
        ex, ey = end
        draw.polygon([(ex, ey), (ex - 18, ey - 10), (ex - 18, ey + 10)], fill="#2e74b5")
    img.save(FLOWCHART_PATH)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10


def add_cover(doc):
    add_para(doc, "LAPORAN UJIAN TENGAH SEMESTER", size=13, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(
        doc,
        "PROYEK DATA MINING",
        size=16,
        bold=True,
        color=DARK_BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=38,
    )
    title = (
        "Optimasi Model Decision Tree untuk Klasifikasi Kondisi Kesehatan "
        "Baterai Lithium-Ion Berdasarkan NASA Battery Aging Dataset"
    )
    add_para(doc, title, size=18, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)

    add_table(
        doc,
        ["Keterangan", "Isi"],
        [
            ("Nama", "[Isi Nama Mahasiswa]"),
            ("NIM", "[Isi NIM]"),
            ("Kelas", "[Isi Kelas]"),
            ("Dosen Pengampu", "[Isi Nama Dosen]"),
            ("Program Studi", "[Isi Program Studi]"),
            ("Universitas", "Universitas Amikom Yogyakarta"),
        ],
        [1.8, 4.7],
    )
    add_para(doc, "Yogyakarta", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_para(doc, "2026", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def build_report():
    create_flowchart()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "Abstrak", 1)
    add_para(
        doc,
        "Penelitian ini membangun model klasifikasi kondisi kesehatan baterai lithium-ion menggunakan algoritma Decision Tree. Dataset yang digunakan berasal dari NASA Battery Aging Dataset, kemudian diolah menjadi data tabular berdasarkan proses discharge. Target klasifikasi terdiri dari tiga kelas, yaitu Healthy, Warning, dan End_of_Life. Evaluasi dilakukan menggunakan nested StratifiedGroupKFold berdasarkan battery_id untuk mengurangi risiko data leakage dan memastikan performa model lebih realistis terhadap baterai yang belum pernah dilihat.",
    )
    add_callout(
        doc,
        "Ringkasan Hasil",
        "Model Decision Tree memperoleh akurasi rata-rata 0,8446 dan F1-score berbobot rata-rata 0,8491. Model akhir memakai criterion gini, max_depth 6, min_samples_leaf 1, dan min_samples_split 5.",
    )

    add_heading(doc, "Daftar Isi", 1)
    for item in [
        "1. Pendahuluan",
        "2. Dataset dan Target Prediksi",
        "3. Metodologi Penelitian",
        "4. Implementasi Program",
        "5. Hasil Evaluasi Model",
        "6. Deployment Aplikasi Web",
        "7. Kesimpulan dan Saran",
        "8. Lampiran",
    ]:
        add_para(doc, item, after=2)
    doc.add_page_break()

    add_heading(doc, "1. Pendahuluan", 1)
    add_heading(doc, "1.1 Latar Belakang", 2)
    add_para(
        doc,
        "Baterai lithium-ion banyak digunakan pada perangkat elektronik, kendaraan listrik, perangkat industri, dan sistem penyimpanan energi. Seiring penggunaan, kapasitas baterai mengalami penurunan sehingga diperlukan metode pemantauan kondisi baterai yang lebih sistematis. Jika kondisi baterai tidak diketahui dengan baik, perangkat dapat mengalami penurunan performa, gangguan operasional, dan risiko kerusakan.",
    )
    add_para(
        doc,
        "Pendekatan data mining dapat digunakan untuk mengidentifikasi pola dari data sensor baterai. Pada proyek ini, model Decision Tree digunakan karena memiliki keunggulan pada sisi interpretasi. Struktur pohon keputusan dapat dijelaskan melalui aturan, sehingga hasil model lebih mudah dianalisis dalam laporan akademik.",
    )
    add_heading(doc, "1.2 Rumusan Masalah", 2)
    for text in [
        "Bagaimana membentuk dataset fitur dari data eksperimen baterai lithium-ion?",
        "Bagaimana membangun model Decision Tree yang akurat dan tidak terlalu overfitting?",
        "Bagaimana mengevaluasi model secara lebih realistis berdasarkan kelompok battery_id?",
        "Bagaimana menyajikan model dalam bentuk aplikasi web sederhana menggunakan Streamlit?",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "1.3 Tujuan Penelitian", 2)
    for text in [
        "Membuat model klasifikasi kesehatan baterai berbasis Decision Tree.",
        "Melakukan EDA, preprocessing, feature engineering, modelling, dan evaluasi secara terstruktur.",
        "Menganalisis performa model menggunakan akurasi, precision, recall, F1-score, dan confusion matrix.",
        "Membangun aplikasi web untuk input sensor dan prediksi kondisi baterai.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "2. Dataset dan Target Prediksi", 1)
    add_para(
        doc,
        "Dataset asli yang digunakan adalah NASA Battery Data Set. Dataset tersebut dapat diunduh melalui tautan berikut: https://phm-datasets.s3.amazonaws.com/NASA/5.+Battery+Data+Set.zip. Data mentah berformat MATLAB kemudian diekstraksi menjadi file data/battery_features.csv.",
    )
    add_table(
        doc,
        ["Komponen", "Keterangan"],
        [
            ("Sumber data", "NASA Battery Data Set"),
            ("Jumlah data", "2.794 baris"),
            ("Jumlah baterai unik", "34 baterai"),
            ("Model utama", "Decision Tree Classifier"),
            ("Target prediksi", "health_status"),
            ("Aplikasi deployment", "Streamlit"),
        ],
        [2.1, 4.4],
    )
    add_heading(doc, "2.1 Kelas Target", 2)
    add_table(
        doc,
        ["Kelas", "Definisi"],
        [
            ("Healthy", "SOH lebih besar atau sama dengan 80%"),
            ("Warning", "SOH berada pada rentang 70% sampai kurang dari 80%"),
            ("End_of_Life", "SOH kurang dari 70%"),
        ],
        [1.8, 4.7],
    )
    add_para(
        doc,
        "Kolom capacity_ahr dan soh tidak digunakan sebagai fitur input karena keduanya berhubungan langsung dengan pembentukan target. Kolom elapsed_time_sec, cycle_index, dan discharge_index juga dikeluarkan agar model tidak mengambil shortcut dari urutan waktu atau siklus.",
    )

    add_heading(doc, "3. Metodologi Penelitian", 1)
    add_picture_if_exists(doc, FLOWCHART_PATH, 6.3, "Gambar 1. Diagram alur penelitian proyek data mining")
    add_heading(doc, "3.1 Tahapan Penelitian", 2)
    for text in [
        "Pengumpulan dataset NASA Battery Data Set.",
        "Ekstraksi data discharge dan pembentukan dataset fitur.",
        "Exploratory Data Analysis untuk memahami distribusi kelas dan pola SOH.",
        "Preprocessing data, penanganan missing value, dan outlier capping.",
        "Feature engineering berbasis tegangan, arus, suhu, dan stabilitas sensor.",
        "Training Decision Tree dan hyperparameter tuning menggunakan GridSearchCV.",
        "Evaluasi model menggunakan nested StratifiedGroupKFold berdasarkan battery_id.",
        "Penyimpanan model akhir dan implementasi dashboard Streamlit.",
    ]:
        add_number(doc, text)

    add_heading(doc, "4. Implementasi Program", 1)
    add_heading(doc, "4.1 Pengambilan Data", 2)
    add_para(doc, "Kode utama pemodelan berada pada file train_decision_tree.py. Dataset fitur dibaca dari data/battery_features.csv.")
    add_code_block(
        doc,
        """
df = pd.read_csv(DATA_PATH)
df = df.drop_duplicates()
df = df.dropna(subset=[TARGET, GROUP_COLUMN])

drop_columns = [
    GROUP_COLUMN, TARGET, "capacity_ahr", "soh",
    "elapsed_time_sec", "cycle_index", "discharge_index",
]
feature_columns = [col for col in df.columns if col not in drop_columns]
X = df[feature_columns]
y = df[TARGET]
groups = df[GROUP_COLUMN]
        """,
    )
    add_heading(doc, "4.2 EDA", 2)
    add_para(
        doc,
        "EDA dilakukan untuk melihat distribusi kelas, tren SOH, dan hubungan antar fitur numerik. Visualisasi ini membantu memastikan bahwa model dibangun berdasarkan pemahaman data, bukan langsung melakukan training tanpa pemeriksaan awal.",
    )
    add_picture_if_exists(doc, BASE_DIR / "outputs" / "01_eda" / "eda_class_distribution.png", 5.2, "Gambar 2. Distribusi kelas health_status")
    add_picture_if_exists(doc, BASE_DIR / "outputs" / "01_eda" / "eda_soh_trend.png", 5.6, "Gambar 3. Tren State of Health pada data baterai")

    add_heading(doc, "4.3 Preprocessing dan Feature Engineering", 2)
    add_para(
        doc,
        "Preprocessing dilakukan dengan median imputer untuk menangani missing value dan IQR capping untuk membatasi pengaruh outlier. Scaling tidak digunakan sebagai tahap utama karena Decision Tree tidak bergantung pada skala fitur.",
    )
    add_table(
        doc,
        ["Kelompok Fitur", "Contoh Fitur", "Alasan"],
        [
            ("Tegangan", "voltage_mean, voltage_min, voltage_std", "Menggambarkan pola tegangan selama discharge"),
            ("Arus", "current_mean, current_min, current_std", "Menggambarkan pola beban dan kestabilan arus"),
            ("Suhu", "temperature_mean, temperature_min, temperature_std", "Menggambarkan respons termal baterai"),
            ("Lingkungan", "ambient_temperature", "Menunjukkan kondisi eksperimen"),
        ],
        [1.45, 2.35, 2.7],
    )
    add_code_block(
        doc,
        """
pipeline = Pipeline(
    steps=[
        ("feature_engineering", BatteryFeatureEngineer()),
        ("outlier_capping", IQRClipper(factor=1.5)),
        ("imputer", SimpleImputer(strategy="median")),
        ("model", DecisionTreeClassifier(random_state=42, class_weight="balanced")),
    ]
)
        """,
    )

    add_heading(doc, "4.4 Modelling dan Validasi", 2)
    add_para(
        doc,
        "Model yang digunakan hanya Decision Tree. Validasi menggunakan nested StratifiedGroupKFold berdasarkan battery_id agar data dari baterai yang sama tidak bercampur secara tidak tepat antara data latih dan data uji.",
    )
    add_code_block(
        doc,
        """
outer_cv = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=42)
inner_cv = StratifiedGroupKFold(n_splits=4, shuffle=True, random_state=42)

grid = GridSearchCV(
    estimator=pipeline,
    param_grid=param_grid,
    scoring="f1_weighted",
    cv=inner_cv,
)
        """,
    )
    add_table(
        doc,
        ["Parameter", "Nilai yang Diuji"],
        [
            ("criterion", "gini, entropy"),
            ("max_depth", "3, 4, 5, 6, 7, 8"),
            ("min_samples_split", "2, 5, 10"),
            ("min_samples_leaf", "1, 3, 5, 10"),
        ],
        [2.1, 4.4],
    )

    add_heading(doc, "5. Hasil Evaluasi Model", 1)
    add_table(
        doc,
        ["Metrik", "Nilai Rata-rata"],
        [
            ("Akurasi", "0,8446 +/- 0,0408"),
            ("F1-score berbobot", "0,8491 +/- 0,0445"),
            ("Precision berbobot", "0,8712 +/- 0,0558"),
            ("Recall berbobot", "0,8446 +/- 0,0408"),
            ("Selisih F1 train-test", "0,0976 +/- 0,0539"),
        ],
        [2.5, 4.0],
    )
    add_para(
        doc,
        "Parameter terbaik model akhir adalah criterion gini, max_depth 6, min_samples_leaf 1, dan min_samples_split 5. Nilai F1-score berbobot menunjukkan bahwa model memiliki performa cukup baik untuk klasifikasi tiga kelas.",
    )
    add_picture_if_exists(doc, BASE_DIR / "outputs" / "04_evaluation" / "confusion_matrix.png", 4.9, "Gambar 4. Confusion matrix hasil evaluasi")
    add_picture_if_exists(doc, BASE_DIR / "outputs" / "04_evaluation" / "feature_importance.png", 5.8, "Gambar 5. Feature importance model Decision Tree")
    add_para(
        doc,
        "Kelas Warning menjadi kelas paling menantang karena berada pada area transisi antara baterai sehat dan baterai yang mendekati akhir masa pakai. Hal ini terlihat dari F1-score kelas Warning sebesar 0,66, lebih rendah dibanding Healthy dan End_of_Life.",
    )

    add_heading(doc, "6. Deployment Aplikasi Web", 1)
    add_para(
        doc,
        "Deployment dibuat menggunakan Streamlit melalui file app.py. Aplikasi menampilkan dashboard ringkasan model, form input sensor pada sidebar, hasil prediksi kelas kesehatan baterai, probabilitas kelas, feature importance, dan performa validasi per fold.",
    )
    add_table(
        doc,
        ["Komponen Deployment", "Keterangan"],
        [
            ("Framework", "Streamlit"),
            ("File utama", "app.py"),
            ("Model", "models/decision_tree_battery.pkl"),
            ("Input pengguna", "Nilai sensor tegangan, arus, suhu, dan suhu lingkungan"),
            ("Output aplikasi", "Kelas Healthy, Warning, atau End_of_Life beserta confidence"),
            ("Cara menjalankan", "streamlit run app.py"),
        ],
        [2.25, 4.25],
    )
    add_callout(
        doc,
        "Catatan Lampiran Screenshot",
        "Setelah aplikasi dijalankan, ambil screenshot halaman dashboard Streamlit dan tempelkan pada bagian ini atau pada lampiran deployment. Screenshot tersebut menjadi bukti bahwa model sudah dapat digunakan melalui antarmuka web.",
    )
    add_code_block(
        doc,
        """
streamlit run app.py
        """,
    )

    add_heading(doc, "7. Kesimpulan dan Saran", 1)
    add_heading(doc, "7.1 Kesimpulan", 2)
    for text in [
        "Model Decision Tree berhasil digunakan untuk mengklasifikasikan kondisi kesehatan baterai lithium-ion menjadi Healthy, Warning, dan End_of_Life.",
        "Evaluasi nested StratifiedGroupKFold berdasarkan battery_id membuat pengujian lebih realistis dan mengurangi risiko data leakage.",
        "Model memperoleh F1-score berbobot rata-rata 0,8491, sehingga cukup layak sebagai model utama dalam proyek ini.",
        "Aplikasi Streamlit berhasil dibuat sebagai antarmuka web untuk menjalankan prediksi berdasarkan input sensor.",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "7.2 Saran", 2)
    for text in [
        "Penelitian berikutnya dapat menambah data baterai atau menguji data dari sumber lain agar generalisasi model lebih kuat.",
        "Kelas Warning perlu dianalisis lebih lanjut karena berada pada area transisi dan paling sulit dibedakan.",
        "Deployment dapat dilanjutkan ke platform cloud agar aplikasi dapat diakses tanpa menjalankan server lokal.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "8. Lampiran", 1)
    add_para(doc, "Lampiran file proyek:")
    for text in [
        "train_decision_tree.py: kode training, tuning, dan evaluasi model.",
        "app.py: kode aplikasi web Streamlit.",
        "data/battery_features.csv: dataset fitur hasil ekstraksi.",
        "models/decision_tree_battery.pkl: model akhir.",
        "outputs/01_eda: hasil exploratory data analysis.",
        "outputs/03_modeling dan outputs/04_evaluation: hasil model dan evaluasi.",
    ]:
        add_bullet(doc, text)
    add_picture_if_exists(doc, BASE_DIR / "outputs" / "03_modeling" / "decision_tree.png", 6.2, "Gambar 6. Struktur pohon keputusan model akhir")

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_report()
    print(path)
